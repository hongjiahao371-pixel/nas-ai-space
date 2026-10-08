"""Native document downloads generated locally from authorized artifact versions."""
from __future__ import annotations
import hashlib
import html
import re
import threading
from pathlib import Path

_LOCK=threading.Lock()
MIMES={'docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
       'pptx':'application/vnd.openxmlformats-officedocument.presentationml.presentation','pdf':'application/pdf'}


def blocks(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].rstrip();i+=1
        if not line.strip():continue
        if line.startswith('```'):
            code=[]
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            i+=1;yield ('code','\n'.join(code));continue
        if line.startswith('|') and i<len(lines) and re.match(r'^\s*\|?[\s:|-]+\|\s*$',lines[i]):
            rows=[[cell.strip() for cell in line.strip('|').split('|')]];i+=1
            while i<len(lines) and lines[i].startswith('|'):rows.append([cell.strip() for cell in lines[i].strip('|').split('|')]);i+=1
            yield ('table',rows);continue
        heading=re.match(r'^(#{1,6})\s+(.*)',line)
        if heading:yield ('h'+str(len(heading[1])),heading[2]);continue
        if re.match(r'^\s*([-*+]\s+|\d+[.)]\s+)',line):yield ('list',re.sub(r'^\s*([-*+]\s+|\d+[.)]\s+)','',line));continue
        yield ('p',line.lstrip('> '))


def plain(value):
    value=re.sub(r'(\*\*|__)(.+?)\1',r'\2',value)
    return re.sub(r'`([^`]+)`',r'\1',value)


def export_artifact(source:Path,title:str,fmt:str,cache:Path):
    if fmt not in MIMES:raise ValueError('请选择 Word、PDF 或 PowerPoint 格式')
    raw=source.read_text(encoding='utf-8')
    key=hashlib.sha256((title+'\0'+raw+'\0'+fmt+'\0v2').encode()).hexdigest()
    cache.mkdir(parents=True,exist_ok=True)
    path=cache/(key+'.'+fmt)
    with _LOCK:
        if path.is_file():return path,MIMES[fmt]
        partial=cache/(key+'.part.'+fmt)
        try:
            data=list(blocks(raw))
            if fmt=='docx':
                from docx import Document
                from docx.shared import Pt,Cm
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn
                doc=Document();doc.core_properties.title=title
                section=doc.sections[0];section.top_margin=Cm(2);section.bottom_margin=Cm(2)
                style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(11)
                fonts=style.element.get_or_add_rPr();east=OxmlElement('w:rFonts');east.set(qn('w:eastAsia'),'Microsoft YaHei');fonts.append(east)
                doc.add_heading(title,0)
                for kind,value in data:
                    if kind=='table':
                        width=max(map(len,value));table=doc.add_table(rows=0,cols=width);table.style='Light Shading Accent 1'
                        for row in value:
                            cells=table.add_row().cells
                            for idx,cell in enumerate(row):cells[idx].text=plain(cell)
                    elif kind.startswith('h'):doc.add_heading(plain(value),min(4,int(kind[1])))
                    elif kind=='list':doc.add_paragraph(plain(value),style='List Bullet')
                    elif kind=='code':
                        p=doc.add_paragraph();run=p.add_run(value);run.font.name='Courier New';run.font.size=Pt(9)
                    else:doc.add_paragraph(plain(value))
                doc.save(partial)
            elif fmt=='pdf':
                from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Preformatted
                from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
                from reportlab.lib.colors import HexColor
                from reportlab.pdfbase import pdfmetrics
                from reportlab.pdfbase.cidfonts import UnicodeCIDFont
                from reportlab.pdfbase.ttfonts import TTFont
                font='STSong-Light'
                for candidate in ['/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc','/System/Library/Fonts/STHeiti Light.ttc']:
                    if Path(candidate).is_file():
                        try:pdfmetrics.registerFont(TTFont('NASChinese',candidate,subfontIndex=0));font='NASChinese';break
                        except Exception:pass
                if font=='STSong-Light':pdfmetrics.registerFont(UnicodeCIDFont(font))
                styles=getSampleStyleSheet()
                for style in styles.byName.values():style.fontName=font;style.wordWrap='CJK'
                styles['Normal'].fontSize=10.5;styles['Normal'].leading=17
                styles['Title'].fontSize=23;styles['Title'].leading=30
                story=[Paragraph(html.escape(title),styles['Title']),Spacer(1,16)]
                def para(v,style='Normal'):return Paragraph(html.escape(plain(v)).replace('\n','<br/>'),styles[style])
                for kind,value in data:
                    if kind=='table':
                        width=max(map(len,value));rows=[[para(cell) for cell in row]+['']*(width-len(row)) for row in value]
                        table=Table(rows,colWidths=[(595-96)/width]*width,repeatRows=1,hAlign='LEFT')
                        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#eef3ff')),('GRID',(0,0),(-1,-1),.4,HexColor('#d8deea')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]));story.extend([table,Spacer(1,10)])
                    else:
                        style='Heading'+str(min(3,int(kind[1]))) if kind.startswith('h') else 'Normal'
                        story.extend([para(('• ' if kind=='list' else '')+value,style),Spacer(1,6)])
                def footer(canvas,doc):
                    canvas.setFont(font,9);canvas.setFillColor(HexColor('#64748b'));canvas.drawRightString(547,28,str(doc.page))
                SimpleDocTemplate(str(partial),pagesize=(595,842),rightMargin=48,leftMargin=48,topMargin=45,bottomMargin=45,title=title,author='NAS AI Space').build(story,onFirstPage=footer,onLaterPages=footer)
            else:
                from pptx import Presentation
                from pptx.util import Inches,Pt
                from pptx.dml.color import RGBColor
                prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
                chunks=[(title,[plain(data[0][1])] if data and data[0][0]=='h1' and data[0][1]!=title else [])];heading=title;paragraphs=[];length=0
                if data and data[0][0]=='h1':data=data[1:]
                def flush():
                    nonlocal paragraphs,length
                    if paragraphs:chunks.append((heading,paragraphs));paragraphs=[];length=0
                for kind,value in data:
                    if kind.startswith('h'):
                        flush();heading=plain(value);continue
                    text='\n'.join(' | '.join(plain(x) for x in row) for row in value) if kind=='table' else plain(value)
                    if kind=='list':text='• '+text
                    # Split long paragraphs as well as long sections; never silently cut content.
                    for start in range(0,len(text),280):
                        item=text[start:start+280]
                        if length+len(item)>550 or len(paragraphs)>=6:flush()
                        paragraphs.append(item);length+=len(item)
                flush()
                if not chunks:chunks=[(title,[''])]
                for number,(heading,paragraphs) in enumerate(chunks,1):
                    slide=prs.slides.add_slide(prs.slide_layouts[6])
                    slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(247,249,252)
                    box=slide.shapes.add_textbox(Inches(.7),Inches(1.8 if number==1 else .45),Inches(11.9),Inches(3.3 if number==1 else 1.0))
                    tf=box.text_frame;tf.word_wrap=True;p=tf.paragraphs[0];p.text=heading;p.font.size=Pt(34 if number==1 else 26);p.font.name='Microsoft YaHei';p.font.color.rgb=RGBColor(25,50,85)
                    body=slide.shapes.add_textbox(Inches(.75),Inches(5.6 if number==1 else 1.65),Inches(11.8),Inches(1.0 if number==1 else 5.1)).text_frame;body.word_wrap=True
                    for idx,text in enumerate(paragraphs):
                        p=body.paragraphs[0] if idx==0 else body.add_paragraph();p.text=text;p.font.name='Microsoft YaHei';p.font.size=Pt(17);p.space_after=Pt(14)
                    foot=slide.shapes.add_textbox(Inches(11.5),Inches(7.0),Inches(1),Inches(.3)).text_frame;foot.text=str(number)
                prs.core_properties.title=title;prs.save(partial)
            partial.chmod(0o600);partial.replace(path)
        finally:partial.unlink(missing_ok=True)
    return path,MIMES[fmt]
