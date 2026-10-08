from __future__ import annotations
import hashlib,json,sqlite3,tarfile,tempfile,unittest,subprocess,shutil,io,time
from pathlib import Path
from dataclasses import replace
from unittest.mock import Mock,patch
from PIL import Image
from app.config import settings
from app.database import Database
from app.services.local_ai import LocalAIClient
from app.services.multimodal import MultimodalService,DIMENSION,validate_embedding
from app.services.search import SearchService
from app.services.scanner import scan_library
from app.services.recovery import RecoveryService
from app.services.artifact_exports import export_artifact
from app.services.media_exports import export_clip,scene_frames
from app.services.notifications import NotificationService
from app.services.extractors import prepare_caption_upgrade_image

class ImprovementsTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.library=self.root/'library';self.library.mkdir();self.data=self.root/'data';self.data.mkdir()
  self.cfg=replace(settings,data_dir=self.data,cache_dir=self.data/'cache',database_path=self.data/'source.db',scan_root=self.library,scan_roots=(self.library,),upload_root=self.root/'uploads',ingest_root=self.root/'uploads'/'inbox',recycle_root=self.data/'recycle',mutation_roots=(),vector_backup_dir=self.data/'snapshots',qdrant_url='http://127.0.0.1:1',multimodal_enabled=True,multimodal_base_url='http://127.0.0.1:1',embedding_base_url='',embedding_model='',notification_webhook_url='')
  self.cfg.cache_dir.mkdir();self.db=Database(self.cfg.database_path);self.db.initialize();self.lib=self.db.create_library('test',str(self.library))
  for name in ['IMG_001.jpg','IMG_002.jpg']:Image.new('RGB',(40,30),'red').save(self.library/name)
  scan_library(self.db,self.lib,lambda *_:None,lambda:False)
  self.files=[self.db.get_file(x) for x in self.db.pending_file_ids()]
 def tearDown(self):self.tmp.cleanup()
 def hit(self,file,score=.5):return {'score':score,'payload':{'file_id':file['id'],'library_id':file['library_id'],'mtime_ns':file['mtime_ns'],'size':file['size'],'source_label':'直接图片','content':'直接图片'}}
 def service(self):return MultimodalService(self.cfg)
 def test_visual_score_beats_caption_word_coverage(self):
  a,b=self.files
  self.db.execute('UPDATE files SET ai_caption=? WHERE id=?',('鸭子 头 潜在 水里',b['id']))
  search=SearchService(self.db,LocalAIClient(self.cfg),Mock(search=Mock(return_value=[])))
  search.multimodal=Mock(enabled=True,search=Mock(return_value=[self.hit(a,.768),self.hit(b,.736)]))
  result=search.search('鸭子头潜在水里',precise=True)
  self.assertEqual(result['results'][0]['id'],a['id'])
  self.assertGreater(result['results'][0]['confidence'],.5)
 def test_visual_query_filters_permissions_disabled_libraries_and_stale_revisions(self):
  search=SearchService(self.db,LocalAIClient(self.cfg),Mock())
  hit=self.hit(self.files[0]);self.assertEqual(len(search.visual_results([hit],'image',20,[self.lib['id']])['results']),1)
  self.assertFalse(search.visual_results([hit],'image',20,[])['results'])
  hit['payload']['mtime_ns']=1;self.assertFalse(search.visual_results([hit],'image',20)['results'])
  self.db.execute('UPDATE libraries SET enabled=0 WHERE id=?',(self.lib['id'],))
  self.assertFalse(search.visual_results([self.hit(self.files[0])],'image',20)['results'])
 def test_shared_query_lease_yields_background_work_and_expires(self):
  one,two=self.service(),self.service()
  with patch('app.services.hardware.memory_runtime',return_value={'available_bytes':8*1024**3}):
   with one.interactive_request():self.assertEqual(two.wait_reason(),'interactive')
   self.assertEqual(two.wait_reason(),'')
   with one._state() as db:db.execute("INSERT INTO interactive_requests VALUES ('expired',?)",(time.time()-1,))
   self.assertEqual(two.wait_reason(),'')
 def test_pause_persists_across_service_instances(self):
  self.service().set_policy({'paused':True,'order':'oldest'})
  self.assertEqual(self.service().run_batch()['waiting'],'paused')
  self.assertEqual(self.service().policy()['order'],'oldest')
 def test_schedule_crosses_midnight(self):
  service=self.service();service.set_policy({'mode':'night','start_hour':23,'end_hour':7})
  with patch('app.services.multimodal.datetime') as clock,patch('app.services.hardware.memory_runtime',return_value={'available_bytes':8*1024**3}):
   clock.now.return_value.hour=2;self.assertEqual(service.wait_reason(),'')
   clock.now.return_value.hour=12;self.assertEqual(service.wait_reason(),'scheduled')
 def test_errors_back_off_stop_and_can_be_manually_retried(self):
  service=self.service();service.wait_reason=Mock(return_value='');service.index_file=Mock(side_effect=ValueError('bad source'))
  for attempt in range(3):
   service.run_batch()
   with service._state() as db:db.execute('UPDATE retry_state SET next_retry=0')
  count=service.index_file.call_count;service.run_batch();self.assertEqual(service.index_file.call_count,count)
  self.assertEqual(service.status()['terminal_errors'],2)
  self.assertEqual(service.retry(),2);service.run_batch();self.assertGreater(service.index_file.call_count,count)
 def test_stale_ready_row_does_not_count_as_new_coverage(self):
  service=self.service();file=self.files[0]
  with service._state() as db:db.execute('INSERT INTO media_index VALUES (?,?,?,?,?,?,?)',(file['id'],'old-signature','ready',1,'image','',time.time()))
  self.assertEqual(service.status()['indexed_files'],0);self.assertEqual(service.status()['pending_files'],2)
 def test_image_signatures_survive_video_sampling_upgrade(self):
  file=self.files[0];legacy=["embeddinggemma2-q8-v1",file['mtime_ns'],file['size'],6,12]
  self.assertEqual(self.service().signature(file),hashlib.sha256(json.dumps(legacy).encode()).hexdigest())
 def test_heic_always_converts_before_caption_upgrade(self):
  file={**self.files[0],'path':str(self.library/'small.heic'),'width':40,'height':30,'size':100}
  (self.library/'small.heic').write_bytes(b'heic')
  with patch('app.services.extractors._convert_image') as convert:
   def write(source,target,*_):Image.new('RGB',(40,30)).save(target)
   convert.side_effect=write
   with prepare_caption_upgrade_image(file,self.cfg) as prepared:
    self.assertEqual(prepared.suffix,'.jpg')
    with Image.open(prepared) as image:self.assertEqual(image.format,'JPEG')
   convert.assert_called_once()
 def test_backup_does_not_touch_main_source_and_contains_both_indexes(self):
  recovery=self.make_recovery();before=hashlib.sha256(self.cfg.database_path.read_bytes()).hexdigest()
  result=recovery.create();self.assertEqual(before,hashlib.sha256(self.cfg.database_path.read_bytes()).hexdigest())
  verified=recovery.verify(result['name']);self.assertTrue(verified['verified']);self.assertEqual(len(verified['collections']),2)
  with tarfile.open(recovery._bundle(result['name'])) as archive:self.assertIn('multimodal.db',archive.getnames())
 def make_recovery(self):
  service=self.service();service.heartbeat('running')
  stores=[]
  for index,collection in enumerate([self.cfg.qdrant_collection,self.cfg.multimodal_collection]):
   folder=self.data/('vectors'+str(index));folder.mkdir(exist_ok=True);(folder/'fake.snapshot').write_bytes(b'snapshot'+bytes([index]))
   store=Mock();store.settings=replace(self.cfg,qdrant_collection=collection,vector_backup_dir=folder)
   response=Mock(status_code=200);response.json.return_value={'result':{'points_count':1,'config':{'params':{'vectors':{'size':DIMENSION if index else 1024,'distance':'Cosine'}}}}}
   store._http.return_value.get.return_value=response;store.create_snapshot.return_value={'name':'fake.snapshot'};stores.append(store)
  service.vectors=stores[1];return RecoveryService(self.db,self.cfg,stores[0],service)
 def test_recovery_rejects_modified_bytes(self):
  recovery=self.make_recovery();name=recovery.create()['name'];path=recovery._bundle(name)
  with tarfile.open(path) as archive:entries=[(m,archive.extractfile(m).read()) for m in archive.getmembers()]
  with tarfile.open(path,'w') as archive:
   for m,data in entries:
    if m.name=='vectors-0.snapshot':data=b'corrupt';m.size=len(data)
    archive.addfile(m,io.BytesIO(data))
  with self.assertRaises(ValueError):recovery.verify(name)
 def test_recovery_refuses_live_destination(self):
  recovery=self.make_recovery();name=recovery.create()['name']
  with self.assertRaises(ValueError):recovery.restore_to(name,self.data,self.cfg.qdrant_url)
 def test_native_exports_preserve_chinese_tables_and_final_paragraph(self):
  source=self.root/'source.md';source.write_text('# 项目说明\n\n这是中文文档。\n\n| 项目 | 数量 |\n| --- | --- |\n| 图片 | 123 |\n\n'+('测试段落'*180)+'\n\n最后一行保留。\n')
  docx,_=export_artifact(source,'测试成果','docx',self.data/'exports');pdf,_=export_artifact(source,'测试成果','pdf',self.data/'exports');pptx,_=export_artifact(source,'测试成果','pptx',self.data/'exports')
  from docx import Document
  doc=Document(docx);self.assertIn('最后一行保留。',''.join(p.text for p in doc.paragraphs));self.assertEqual(doc.tables[0].cell(1,1).text,'123')
  from pypdf import PdfReader
  self.assertIn('最后一行保留。',''.join(p.extract_text() for p in PdfReader(pdf).pages))
  from pptx import Presentation
  prs=Presentation(pptx);self.assertGreater(len(prs.slides),1);self.assertIn('最后一行保留。',''.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame))
 def test_unconfigured_notification_has_no_outbound_request(self):
  notify=NotificationService(self.db,self.cfg);notify.enqueue('1','test','title','body')
  with patch('httpx.Client') as client:self.assertFalse(notify.flush()['configured']);client.assert_not_called()
 def test_notification_failure_is_durable_and_deduplicated(self):
  notify=NotificationService(self.db,replace(self.cfg,notification_webhook_url='http://127.0.0.1:1/hook'))
  notify.enqueue('same','test','title','body');notify.enqueue('same','test','title','body')
  notify.flush();rows=self.db.fetchall('SELECT * FROM external_notifications');self.assertEqual(len(rows),1);self.assertEqual(rows[0]['attempts'],1);self.assertEqual(rows[0]['status'],'pending');self.assertNotIn('/hook',rows[0]['error'])
 @unittest.skipUnless(shutil.which('ffmpeg'),'ffmpeg required')
 def test_clip_export_preserves_source_and_validates_duration(self):
  source=self.library/'test.mp4';subprocess.run(['ffmpeg','-hide_banner','-loglevel','error','-f','lavfi','-i','color=c=red:s=128x128:d=3','-c:v','libx264',str(source)],check=True)
  scan_library(self.db,self.lib,lambda *_:None,lambda:False)
  file=self.db.fetchone('SELECT * FROM files WHERE name=?',('test.mp4',));self.db.execute('UPDATE files SET duration=3 WHERE id=?',(file['id'],))
  export_id=self.db.execute("INSERT INTO clip_exports(file_id,start_time,end_time,created_at) VALUES (?,.5,1.5,'test')",(file['id'],));before=hashlib.sha256(source.read_bytes()).hexdigest()
  result=export_clip(self.db,self.cfg,export_id);self.assertAlmostEqual(result['duration'],1,delta=.1);self.assertEqual(before,hashlib.sha256(source.read_bytes()).hexdigest())
  self.assertEqual(self.db.fetchone('SELECT status FROM clip_exports WHERE id=?',(export_id,))['status'],'ready')

class ApiImprovementsTests(unittest.TestCase):
 hit=ImprovementsTests.hit
 def setUp(self):
  ImprovementsTests.setUp(self)
  self.cfg=replace(self.cfg,api_token='t'*48,automatic_backup_enabled=False,auto_index_enabled=False,watch_enabled=False,vision_model='',chat_model='',transcription_base_url='')
  import app.main as main
  self.main=main
  from fastapi.testclient import TestClient
  self.patch=patch.object(main,'settings',self.cfg);self.patch.start();self.client=TestClient(main.app);self.client.__enter__()
  self.headers={'Authorization':'Bearer '+self.cfg.api_token}
 def tearDown(self):self.client.__exit__(None,None,None);self.patch.stop();ImprovementsTests.tearDown(self)
 def test_image_api_crop_validation_permissions_and_invalid_content(self):
  state=self.main.state
  with patch.object(state.search.multimodal,'image_search',return_value=[self.hit(self.files[0])]) as query:
   data=(self.library/'IMG_001.jpg').read_bytes()
   r=self.client.post('/api/search/image?left=.25&top=.25&right=.75&bottom=.75',content=data,headers=self.headers)
   self.assertEqual(r.status_code,200,r.text);self.assertEqual(r.json()['results'][0]['id'],self.files[0]['id']);query.assert_called_once()
   self.assertEqual(self.client.post('/api/search/image?left=.9&right=.1',content=data,headers=self.headers).status_code,400)
   self.assertEqual(self.client.post('/api/search/image',content=b'not an image',headers=self.headers).status_code,400)
   self.assertEqual(self.client.post('/api/search/image',content=data).status_code,401)
 def test_policy_endpoint_rejects_invalid_order_and_preserves_pause(self):
  r=self.client.put('/api/index/multimodal/policy',json={'paused':True,'order':'oldest'},headers=self.headers);self.assertEqual(r.status_code,200,r.text)
  r=self.client.get('/api/index/status',headers=self.headers);self.assertTrue(r.json()['multimodal']['policy']['paused'])
  self.assertEqual(self.client.put('/api/index/multimodal/policy',json={'order':'invalid'},headers=self.headers).status_code,400)
 def test_recovery_endpoints_require_admin_and_reject_paths(self):
  self.assertEqual(self.client.get('/api/operations/recovery').status_code,401)
  self.assertEqual(self.client.post('/api/operations/recovery/not-a-bundle/verify',headers=self.headers).status_code,409)
  state=self.main.state;state.tasks.quiescing=True
  self.assertEqual(self.client.post('/api/index/multimodal/retry',headers=self.headers).status_code,503)
  state.tasks.quiescing=False
 def test_clip_api_rejects_image_and_bad_intervals(self):
  file=self.files[0];r=self.client.post(f"/api/files/{file['id']}/clips",json={'start':0,'end':2},headers=self.headers);self.assertEqual(r.status_code,400)
  self.db.execute("UPDATE files SET kind='video',duration=5 WHERE id=?",(file['id'],))
  self.assertEqual(self.client.post(f"/api/files/{file['id']}/clips",json={'start':3,'end':2},headers=self.headers).status_code,400)
  self.assertEqual(self.client.post(f"/api/files/{file['id']}/clips",json={'start':0,'end':6},headers=self.headers).status_code,400)
  self.assertEqual(self.client.get('/api/clips/999999',headers=self.headers).status_code,404)
 def test_native_downloads_use_existing_artifact_permissions(self):
  from app.security import hash_password
  user=self.db.create_user('export-qa','export-qa',hash_password('TestPassword123'),'admin',[])
  login=self.client.post('/api/auth/login',json={'username':'export-qa','password':'TestPassword123'})
  self.assertEqual(login.status_code,200,login.text)
  headers={'Authorization':'Bearer '+login.json()['token']}
  service=self.main.state.productivity
  artifact=service.create_artifact('原生导出测试','report',user['id'])
  with patch.object(service.ai,'generate_artifact',return_value='# 原生导出测试\n\n中文内容及最后一行。'):
   version=service.generate_artifact_version(artifact['id'],'test',[self.files[0]['id']],user['id'])
  for fmt,signature in [('docx',b'PK'),('pptx',b'PK'),('pdf',b'%PDF')]:
   url=f"/api/artifacts/{artifact['id']}/versions/{version['id']}/download?format={fmt}"
   response=self.client.get(url,headers=headers);self.assertEqual(response.status_code,200,response.text[:200] if response.status_code!=200 else '');self.assertTrue(response.content.startswith(signature));self.assertEqual(self.client.get(url).status_code,401)
